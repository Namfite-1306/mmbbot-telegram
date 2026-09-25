from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Bao_cao_ngan_Mr_Mission_Bossible_Bot.docx"
FONT = "Arial"
NAVY = "1F4E78"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"


def set_cell_fill(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_borders(cell, color=LIGHT_GRAY, size="6"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_run_font(run, size=None, bold=None, color="000000"):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    set_run_font(paragraph.add_run(text), 10.5)
    return paragraph


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_fill(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(paragraph.add_run(text), 9.5, True, "FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column, value in enumerate(values):
            cell = cells[column]
            cell.width = Inches(widths[column])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2:
                set_cell_fill(cell, PALE_BLUE)
            set_cell_borders(cell)
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            set_run_font(paragraph.add_run(value), 9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08

title_style = styles["Title"]
title_style.font.name = FONT
title_style.font.size = Pt(20)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
title_style_ppr = title_style._element.get_or_add_pPr()
title_style_border = title_style_ppr.find(qn("w:pBdr"))
if title_style_border is not None:
    title_style_ppr.remove(title_style_border)

for name, size in (("Heading 1", 14), ("Heading 2", 11.5)):
    style = styles[name]
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.paragraph_format.space_before = Pt(9)
    style.paragraph_format.space_after = Pt(4)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(6)
title_ppr = title._p.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
title.add_run("Báo cáo ngắn về Mr Mission Bossible Bot")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(12)
set_run_font(subtitle.add_run("Chức năng dữ liệu API và cách tính điểm"), 11, True)

intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(8)
set_run_font(intro.add_run(
    "Tài liệu này giúp người mới hiểu bot làm gì, dữ liệu được lấy và lưu ở đâu, "
    "các file code chính nằm ở đâu và điểm T F M được tính như thế nào. Bot chỉ phục vụ "
    "học tập và giao dịch mô phỏng, không gửi lệnh thật đến công ty chứng khoán."), 10.5)

doc.add_heading("1 Bot làm gì", level=1)
for item in (
    "Phân tích một cổ phiếu bằng /soi MÃ hoặc /analyze MÃ.",
    "Quét toàn thị trường bằng /scan từ dữ liệu đã lưu trong SQLite.",
    "Quản lý danh mục quan tâm bằng /add, /remove và /portfolio.",
    "Đề xuất phân bổ vốn mô phỏng bằng /modelportfolio hoặc /phanbo.",
    "Mua bán ảo bằng /paper, /paperbuy và /papersell; không đặt lệnh thật.",
):
    add_bullet(doc, item)

doc.add_heading("2 Luồng dữ liệu", level=1)
flow = doc.add_paragraph()
flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
flow.paragraph_format.space_after = Pt(8)
set_run_font(flow.add_run(
    "DNSE  Vietcap  CafeF   →   Chuẩn hóa và kiểm tra   →   SQLite   →   Tính tín hiệu   →   Telegram"),
    10.2, True)
add_bullet(doc, "Tác vụ nền kiểm tra phiên mới và cập nhật toàn thị trường vào database.")
add_bullet(doc, "/scan chỉ đọc database, không crawl toàn thị trường khi người dùng đang chờ.")
add_bullet(doc, "/soi dùng dữ liệu cuối ngày đã lưu để tính điểm và có thể lấy thêm giá DNSE mới nhất để hiển thị.")

doc.add_heading("3 Các lệnh chính", level=1)
add_table(doc, ["Lệnh", "Công dụng"], [
    ("/soi MÃ  /analyze MÃ", "Phân tích một mã, xem biểu đồ và nút mua bán ảo."),
    ("/scan", "Quét BUY SELL và mã cần theo dõi từ SQLite."),
    ("/add  /remove  /portfolio", "Quản lý và xem danh mục quan tâm."),
    ("/modelportfolio  /phanbo", "Xem cơ cấu phân bổ vốn mô phỏng."),
    ("/paper", "Xem tài khoản, vị thế và biểu đồ tài sản ảo."),
    ("/paperbuy  /buy  /mua", "Đặt lệnh mua ảo với mã và số cổ phiếu."),
    ("/papersell  /sell  /ban", "Đặt lệnh bán ảo với mã và số cổ phiếu."),
    ("/status", "Kiểm tra database và trạng thái cập nhật dữ liệu."),
], [2.2, 4.7])

doc.add_heading("4 Các file code chính", level=1)
add_table(doc, ["File", "Vai trò"], [
    ("app/main.py", "Khởi động bot, menu Telegram và tác vụ nền."),
    ("app/bot/handlers.py", "Nhận lệnh hoặc nút và gọi đúng chức năng."),
    ("app/bot/formatter.py", "Tạo nội dung phản hồi cho người dùng."),
    ("app/market_store.py", "Điều phối tải, kiểm tra và lưu dữ liệu thị trường."),
    ("app/strategy_engine.py", "Tính điểm T F M và điều kiện chiến lược."),
    ("app/providers/strategy_signal_provider.py", "Chuyển kết quả tính toán thành tín hiệu."),
    ("app/paper_broker.py", "Quản lý tiền, vị thế và lệnh giao dịch ảo."),
    ("app/storage/database.py", "Khai báo cấu trúc các bảng SQLite."),
], [2.75, 4.15])

doc.add_heading("5 File gọi API", level=1)
add_table(doc, ["Nguồn", "File gọi API", "Dữ liệu"], [
    ("DNSE", "app/dnse_market_data.py", "OHLCV, chỉ số và giá khớp mới nhất."),
    ("Vietcap", "crawl data/Project_python3/vietcap_common.py", "HTTP client dùng chung, retry và giới hạn tốc độ."),
    ("Vietcap", "crawl data/Project_python3/vietcap_market_data.py", "Danh sách mã, OHLC, intraday và priceboard."),
    ("Vietcap", "crawl data/Project_python3/vietcap_fundamental_data.py", "Báo cáo tài chính và dữ liệu cơ bản."),
    ("CafeF", "app/cafef_market_data.py", "Giá dự phòng khi nguồn chính lỗi."),
    ("CafeF", "app/cafef_news.py", "Tin doanh nghiệp cho lệnh soi."),
], [1.0, 3.35, 2.55])

doc.add_heading("6 Data được lưu ở đâu", level=1)
add_table(doc, ["Loại dữ liệu", "File hoặc thư mục"], [
    ("Database bot đang dùng", "data/fintech_bot.db"),
    ("Giá từng cổ phiếu", "crawl data/Project_python3/data/market_data/<MÃ>.csv"),
    ("Báo cáo tài chính", "crawl data/Project_python3/data/fundamental/<MÃ>_fundamental.csv"),
    ("Chỉ số thị trường", "crawl data/Project_python3/data/index_data/VNINDEX.csv"),
    ("Danh sách mã và sàn", "symbol_info_cache.csv và symbols.csv"),
    ("Bảng giá", "priceboard và priceboard_daily"),
    ("Lỗi và chất lượng dữ liệu", "crawl data/Project_python3/data/quality"),
], [2.1, 4.8])
paragraph = doc.add_paragraph()
set_run_font(paragraph.add_run(
    "CSV dùng để lưu vết dữ liệu thu thập. SQLite là nguồn bot truy cập nhanh khi phục vụ lệnh người dùng."),
    10.2, True)

doc.add_page_break()
doc.add_heading("7 Cách tính điểm ngắn gọn", level=1)
add_table(doc, ["Điểm", "Cách tính"], [
    ("T  Kỹ thuật", "Close trên EMA20: 30; EMA20 trên EMA50: 30; EMA50 trên EMA200: 25; RSI14 từ 50 đến 70: 15. Tối đa 100."),
    ("F  Cơ bản", "Chấm tăng trưởng doanh thu, lợi nhuận, ROE, dòng tiền kinh doanh và nợ trên vốn chủ; sau đó chuẩn hóa về 100. Thiếu dữ liệu quan trọng thì để trống."),
    ("M  Sức mạnh", "So cổ phiếu với chỉ số trong 5 và 20 phiên. Vượt chỉ số ở mỗi khoảng được 50 điểm, nên M thường là 0, 50 hoặc 100."),
    ("Điểm tổng", "0,375 x T + 0,3125 x F + 0,3125 x M. Chỉ tính khi đủ cả T F M."),
], [1.35, 5.55])

paragraph = doc.add_paragraph()
paragraph.paragraph_format.space_before = Pt(2)
set_run_font(paragraph.add_run("Lưu ý  "), 10.5, True)
set_run_font(paragraph.add_run(
    "Điểm cao chưa tự động tạo lệnh. Bot còn kiểm tra xu hướng thị trường, breakout, khối lượng và điều kiện an toàn."),
    10.5)

doc.add_heading("8 Điều cần nhớ", level=1)
for item in (
    "Kiểm tra /status để biết ngày và độ đầy đủ của dữ liệu trước khi dùng /scan.",
    "Nếu một mã không có dữ liệu, bot phải báo rõ thay vì tự coi đó là tín hiệu bán.",
    "Token và API key chỉ lưu trong .env, không ghi vào code hoặc commit lên Git.",
    "Sau khi sửa code, cần khởi động lại bằng python -m app.main.",
):
    add_bullet(doc, item)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(footer.add_run("Mr Mission Bossible Bot  |  Báo cáo cập nhật 25/09/2026"), 8.5, False, "666666")

doc.save(OUTPUT)
print(OUTPUT)
